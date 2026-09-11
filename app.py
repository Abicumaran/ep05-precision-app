import io
import zipfile
from dataclasses import dataclass, field
from typing import List, Tuple, Dict, Optional

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from scipy import stats
import streamlit as st
import statsmodels.formula.api as smf

# -----------------------------
# Defaults that match your EP05 template + current CBC paired dataset
# -----------------------------
REQUIRED_BASE_COLS = ["batch_id", "Blood Sample ID", "Level", "Day", "Replicate", "Device"]
BASE_COL_ALIASES = {
    "Blood Sample ID": ["Blood Sample ID", "bloodSampleId", "blood_sample_id", "sample_id"],
    "Device": ["Device", "deviceId", "device_id", "serialNumber", "serial_number"],
}

# User-specified device/reference pairs.  NEU is kept as the display analyte, while
# NEUT_2/NEUT_ref are accepted because that is how the uploaded dataset labels them.
DEFAULT_ANALYTE_PAIRS = {
    "RBC": ("RBC", "RBC_ref"),
    "WBC": ("WBC_2", "WBC_ref"),
    "PLT": ("PLT", "PLT_ref"),
    "HCT": ("HCT", "HCT_ref"),
    "HGB": ("HGB", "HGB_ref"),
    "MCV": ("MCV", "MCV_ref"),
    "RDW": ("RDW", "RDW_ref"),
    "MCH": ("MCH", "MCH_ref"),
    "MCHC": ("MCHC", "MCHC_ref"),
    "NEU": (["NEU_2", "NEUT_2", "NEU", "NEUT"], ["NEU_ref", "NEUT_ref"]),
    "LYMPH": ("LYMPH_2", "LYMPH_ref"),
    "MXD": ("MXD_2", "MXD_ref"),
}

DEFAULT_ANALYTES = list(DEFAULT_ANALYTE_PAIRS.keys())

# Keep analyte/model choices in the same user-facing order as the Short-Term app.
# These are uploaded device-result column names, not a change to the EP05 method.
DEFAULT_ANALYTE_COLUMN_ORDER = [
    "RBC", "WBC_2", "PLT", "HCT", "HGB", "MCV", "RDW", "MCH", "MCHC",
    "NEUT_2", "LYMPH_2", "MXD_2", "PLT_3", "MCV_3", "RDW_3",
]
DEFAULT_LEVELS = ["Low", "Mid", "High"]
DEFAULT_DAYS = ["D1", "D2", "D3", "D4", "D5"]
DEFAULT_REPLICATES = [1, 2, 3, 4, 5]
DEFAULT_DEVICES = []

NORMALIZATION_METHODS = [
    "Raw/no normalization",
    "Day 1 anchoring: reference only",
    "Day 1 anchoring: device and reference separately",
    "Per-level median centering: reference only",
    "Per-level median centering: device and reference separately",
    "Reference drift correction: day-wise reference factors applied to paired values",
    "Robust median/MAD z-score: reference only",
    "Robust median/MAD z-score: device and reference separately",
]

VALUE_OUTPUT_MODES = [
    "Device normalized value",
    "Reference normalized value",
    "Bias: device - reference",
    "%Bias: 100*(device-reference)/reference",
]

# -----------------------------
# Core config
# -----------------------------
@dataclass
class Config:
    analytes: List[str]
    levels: List[str]
    days: List[str]
    replicates: List[int]
    devices: List[str]
    gcrit: float
    expected_n: int
    device_mode: str
    outlier_method: str
    max_remove_per_group: int
    gcrit_mode: str
    gcrit_alpha: float
    gcrit_tail: str
    modified_z_threshold: float
    robust_interval_z: float
    do_bootstrap_ci: bool
    n_boot: int
    seed: int
    paired_analytes: List[str] = field(default_factory=list)
    normalization_method: str = "Raw/no normalization"
    value_output_modes: List[str] = field(default_factory=lambda: ["Device normalized value"])
    analyte_pair_map: Dict[str, Tuple[str, str]] = field(default_factory=dict)
    global_flag_col: Optional[str] = None
    treat_all_global_false: bool = False


# -----------------------------
# Utility helpers
# -----------------------------
def first_existing_column(df: pd.DataFrame, candidates) -> Optional[str]:
    if isinstance(candidates, str):
        candidates = [candidates]
    for c in candidates:
        if c in df.columns:
            return c
    return None


def standardize_base_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Accept the uploaded dataset aliases without requiring manual column edits."""
    out = df.copy()
    for canonical, aliases in BASE_COL_ALIASES.items():
        if canonical not in out.columns:
            found = first_existing_column(out, aliases)
            if found is not None:
                out[canonical] = out[found]
    if "Device" not in out.columns and "deviceId" in out.columns:
        out["Device"] = out["deviceId"].astype(str)
    return out


def resolve_analyte_pairs(df: pd.DataFrame, selected: Optional[List[str]] = None) -> Dict[str, Tuple[str, str]]:
    selected = selected or list(DEFAULT_ANALYTE_PAIRS.keys())
    pairs = {}
    for label in selected:
        if label not in DEFAULT_ANALYTE_PAIRS:
            continue
        dev_spec, ref_spec = DEFAULT_ANALYTE_PAIRS[label]
        dev_col = first_existing_column(df, dev_spec)
        ref_col = first_existing_column(df, ref_spec)
        if dev_col is not None and ref_col is not None:
            pairs[label] = (dev_col, ref_col)
    return pairs


def normalize_bool(value) -> bool:
    if pd.isna(value):
        return False
    if isinstance(value, (bool, np.bool_)):
        return bool(value)
    if isinstance(value, (int, float, np.integer, np.floating)) and np.isfinite(value):
        return bool(int(value))
    return str(value).strip().lower() in {"true", "t", "1", "yes", "y", "flagged"}


def split_global_flag_rows(df: pd.DataFrame, cfg: Config) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """Exclude QC-rejected rows before normalization/statistics; keep them for a separate audit sheet."""
    if cfg.treat_all_global_false or not cfg.global_flag_col or cfg.global_flag_col == "None":
        return df.copy(), pd.DataFrame()
    if cfg.global_flag_col not in df.columns:
        raise ValueError(f"Selected global flag column '{cfg.global_flag_col}' was not found.")
    mask = df[cfg.global_flag_col].map(normalize_bool).fillna(False).astype(bool)
    return df.loc[~mask].copy(), df.loc[mask].copy()


def default_analyte_mapping_table(df: pd.DataFrame) -> pd.DataFrame:
    """Editable starter map: validated defaults plus additional PLT model columns when present."""
    rows = []
    base_pairs = resolve_analyte_pairs(df)
    for label, (dev_col, ref_col) in base_pairs.items():
        rows.append({"Include": True, "Analyte": label, "Device column": dev_col, "Reference column": ref_col})

    # Additional model columns can share the same validated reference analyte.
    # This only expands user-selectable mappings; it does not alter the EP05 calculations.
    extra_models = [
        ("PLT_2", "PLT 2", "PLT_ref"),
        ("PLT_3", "PLT 3", "PLT_ref"),
        ("MCV_3", "MCV 3", "MCV_ref"),
        ("RDW_3", "RDW 3", "RDW_ref"),
    ]
    for dev_col, label, ref_col in extra_models:
        if dev_col in df.columns and ref_col in df.columns and not any(r["Device column"] == dev_col for r in rows):
            rows.append({"Include": True, "Analyte": label, "Device column": dev_col, "Reference column": ref_col})

    # Reorder starter rows to mirror the Short-Term analyte selector, then append
    # any additional mappings that are not part of that preferred list.
    order_index = {name: i for i, name in enumerate(DEFAULT_ANALYTE_COLUMN_ORDER)}
    rows.sort(key=lambda r: (order_index.get(str(r["Device column"]), len(order_index)), str(r["Device column"])))
    return pd.DataFrame(rows, columns=["Include", "Analyte", "Device column", "Reference column"])


def parse_analyte_mapping_table(table: pd.DataFrame, df: pd.DataFrame) -> Dict[str, Tuple[str, str]]:
    mapping = {}
    if table is None or table.empty:
        return mapping
    for _, r in table.iterrows():
        include = r.get("Include", True)
        if pd.isna(include) or not bool(include):
            continue
        label = str(r.get("Analyte", "")).strip()
        dev_col = str(r.get("Device column", "")).strip()
        ref_col = str(r.get("Reference column", "")).strip()
        if not label or not dev_col or not ref_col or dev_col == "nan" or ref_col == "nan":
            continue
        if dev_col not in df.columns or ref_col not in df.columns:
            continue
        mapping[label] = (dev_col, ref_col)
    return mapping


def safe_divide(num, den):
    den = pd.to_numeric(den, errors="coerce")
    num = pd.to_numeric(num, errors="coerce")
    return num / den.replace(0, np.nan)


def robust_sd_mad(x: np.ndarray) -> float:
    x = np.asarray(x, dtype=float)
    x = x[np.isfinite(x)]
    if len(x) == 0:
        return np.nan
    med = np.median(x)
    mad = np.median(np.abs(x - med))
    return 1.4826 * mad


def robust_mad_series(s: pd.Series) -> float:
    vals = pd.to_numeric(s, errors="coerce").dropna().to_numpy(dtype=float)
    return robust_sd_mad(vals)


def iqr(x: np.ndarray) -> float:
    x = np.asarray(x, dtype=float)
    x = x[np.isfinite(x)]
    if len(x) == 0:
        return np.nan
    return float(np.percentile(x, 75) - np.percentile(x, 25))


def bootstrap_ci_mean_or_median(y: np.ndarray, use_median: bool, n_boot: int, seed: int) -> Tuple[float, float]:
    y = np.asarray(y, dtype=float)
    y = y[np.isfinite(y)]
    if len(y) < 2:
        return np.nan, np.nan
    rng = np.random.default_rng(int(seed))
    vals = []
    for _ in range(int(max(200, n_boot))):
        b = rng.choice(y, size=len(y), replace=True)
        vals.append(np.median(b) if use_median else np.mean(b))
    return tuple(np.percentile(vals, [2.5, 97.5]).astype(float))


def grubbs_gcrit_auto(n: int, alpha: float = 0.01, tail: str = "Two-sided") -> float:
    n = int(n)
    if n < 3:
        return np.nan
    alpha = float(alpha)
    if str(tail).lower().startswith("one"):
        t_quant = 1.0 - alpha / n
    else:
        t_quant = 1.0 - alpha / (2.0 * n)
    t = stats.t.ppf(t_quant, df=n - 2)
    return float(((n - 1) / np.sqrt(n)) * np.sqrt((t * t) / (n - 2 + t * t)))


def current_gcrit(n: int, cfg: Config) -> float:
    if str(cfg.gcrit_mode).startswith("Automatic"):
        return grubbs_gcrit_auto(n=n, alpha=cfg.gcrit_alpha, tail=cfg.gcrit_tail)
    return float(cfg.gcrit)


# -----------------------------
# Normalization and paired metrics
# -----------------------------

def day1_anchor_by_level(df: pd.DataFrame, values: pd.Series) -> pd.Series:
    tmp = df.assign(_x=pd.to_numeric(values, errors="coerce"))
    day_mask = tmp["Day"].astype(str).isin(["D1", "1", "Day1"])
    anchors = tmp.loc[day_mask].groupby("Level")["_x"].median()
    fallback = tmp.groupby("Level")["_x"].median()
    anchors = fallback.combine_first(anchors) if anchors.empty else anchors.combine_first(fallback)
    return tmp["Level"].map(anchors).astype(float)

def group_day_medians(df: pd.DataFrame, col: str, group_cols: List[str]) -> pd.Series:
    return df.groupby(group_cols)[col].transform(lambda x: pd.to_numeric(x, errors="coerce").median())


def apply_normalization_for_pair(df: pd.DataFrame, label: str, dev_col: str, ref_col: str, method: str) -> pd.DataFrame:
    out = df.copy()
    dev = pd.to_numeric(out[dev_col], errors="coerce")
    ref = pd.to_numeric(out[ref_col], errors="coerce")
    norm_dev = dev.copy()
    norm_ref = ref.copy()

    # Anchors are estimated within Level because Low/Mid/High are true concentration strata.
    level_cols = ["Level"]
    level_day_cols = ["Level", "Day"]

    if method == "Day 1 anchoring: reference only":
        day_center = group_day_medians(out.assign(_ref=ref), "_ref", level_day_cols)
        anchor = day1_anchor_by_level(out, ref)
        norm_ref = ref * safe_divide(anchor, day_center)

    elif method == "Day 1 anchoring: device and reference separately":
        for raw, name in [(dev, "dev"), (ref, "ref")]:
            tmp = out.assign(_x=raw)
            day_center = group_day_medians(tmp, "_x", level_day_cols)
            anchor = day1_anchor_by_level(out, raw)
            adj = raw * safe_divide(anchor, day_center)
            if name == "dev":
                norm_dev = adj
            else:
                norm_ref = adj

    elif method == "Per-level median centering: reference only":
        level_center = group_day_medians(out.assign(_ref=ref), "_ref", level_cols)
        norm_ref = ref - level_center

    elif method == "Per-level median centering: device and reference separately":
        dev_center = group_day_medians(out.assign(_dev=dev), "_dev", level_cols)
        ref_center = group_day_medians(out.assign(_ref=ref), "_ref", level_cols)
        norm_dev = dev - dev_center
        norm_ref = ref - ref_center

    elif method == "Reference drift correction: day-wise reference factors applied to paired values":
        tmp = out.assign(_ref=ref)
        day_center = group_day_medians(tmp, "_ref", level_day_cols)
        anchor = day1_anchor_by_level(out, ref)
        factor = safe_divide(anchor, day_center)
        norm_dev = dev * factor
        norm_ref = ref * factor

    elif method == "Robust median/MAD z-score: reference only":
        med = group_day_medians(out.assign(_ref=ref), "_ref", level_cols)
        mad_sd = out.assign(_ref=ref).groupby("Level")["_ref"].transform(robust_mad_series).replace(0, np.nan)
        norm_ref = (ref - med) / mad_sd

    elif method == "Robust median/MAD z-score: device and reference separately":
        dev_med = group_day_medians(out.assign(_dev=dev), "_dev", level_cols)
        ref_med = group_day_medians(out.assign(_ref=ref), "_ref", level_cols)
        dev_mad = out.assign(_dev=dev).groupby("Level")["_dev"].transform(robust_mad_series).replace(0, np.nan)
        ref_mad = out.assign(_ref=ref).groupby("Level")["_ref"].transform(robust_mad_series).replace(0, np.nan)
        norm_dev = (dev - dev_med) / dev_mad
        norm_ref = (ref - ref_med) / ref_mad

    out[f"{label}__device_norm"] = norm_dev
    out[f"{label}__ref_norm"] = norm_ref
    out[f"{label}__bias"] = norm_dev - norm_ref
    out[f"{label}__pctbias"] = 100.0 * safe_divide(norm_dev - norm_ref, norm_ref)
    return out


def build_analysis_dataframe(df: pd.DataFrame, cfg: Config) -> Tuple[pd.DataFrame, List[str], pd.DataFrame]:
    work = standardize_base_columns(df)
    analysis_cols = []
    pair_rows = []
    for label, (dev_col, ref_col) in cfg.analyte_pair_map.items():
        work = apply_normalization_for_pair(work, label, dev_col, ref_col, cfg.normalization_method)
        pair_rows.append({"analyte": label, "device_column": dev_col, "reference_column": ref_col})
        for mode in cfg.value_output_modes:
            if mode == "Device normalized value":
                analysis_cols.append(f"{label}__device_norm")
            elif mode == "Reference normalized value":
                analysis_cols.append(f"{label}__ref_norm")
            elif mode == "Bias: device - reference":
                analysis_cols.append(f"{label}__bias")
            elif mode == "%Bias: 100*(device-reference)/reference":
                analysis_cols.append(f"{label}__pctbias")
    # Preserve manually selected raw columns, if any.
    for a in cfg.analytes:
        if a in work.columns and a not in analysis_cols:
            analysis_cols.append(a)
    return work, analysis_cols, pd.DataFrame(pair_rows)


def evaluate_normalization_methods(df: pd.DataFrame, analyte_pair_map: Dict[str, Tuple[str, str]], methods: List[str]) -> pd.DataFrame:
    rows = []
    if not analyte_pair_map:
        return pd.DataFrame()
    base = standardize_base_columns(df)
    for method in methods:
        tmp = base.copy()
        for label, (dev_col, ref_col) in analyte_pair_map.items():
            tmp = apply_normalization_for_pair(tmp, label, dev_col, ref_col, method)
            for level in sorted(tmp["Level"].dropna().astype(str).unique()):
                sub = tmp[tmp["Level"].astype(str) == level]
                if sub.empty:
                    continue
                ref_day_meds = sub.groupby("Day")[f"{label}__ref_norm"].median(numeric_only=True).to_numpy(dtype=float)
                dev_day_meds = sub.groupby("Day")[f"{label}__device_norm"].median(numeric_only=True).to_numpy(dtype=float)
                pctbias = pd.to_numeric(sub[f"{label}__pctbias"], errors="coerce").dropna().to_numpy(dtype=float)
                ref_cv = 100.0 * np.nanstd(ref_day_meds, ddof=1) / np.nanmean(np.abs(ref_day_meds)) if len(ref_day_meds) > 1 and np.nanmean(np.abs(ref_day_meds)) != 0 else np.nan
                dev_cv = 100.0 * np.nanstd(dev_day_meds, ddof=1) / np.nanmean(np.abs(dev_day_meds)) if len(dev_day_meds) > 1 and np.nanmean(np.abs(dev_day_meds)) != 0 else np.nan
                rows.append({
                    "normalization_method": method,
                    "analyte": label,
                    "Level": level,
                    "ref_day_median_CV_%": ref_cv,
                    "device_day_median_CV_%": dev_cv,
                    "median_%bias": float(np.nanmedian(pctbias)) if len(pctbias) else np.nan,
                    "abs_median_%bias": float(abs(np.nanmedian(pctbias))) if len(pctbias) else np.nan,
                    "IQR_%bias": iqr(pctbias),
                    "N": int(len(sub)),
                })

        # Level separation preservation is computed per analyte across all levels.
        for label in analyte_pair_map:
            if f"{label}__ref_norm" not in tmp.columns:
                continue
            centers = tmp.groupby("Level")[f"{label}__ref_norm"].median(numeric_only=True)
            if all(l in centers.index for l in ["Low", "Mid", "High"]):
                sep_min = min(abs(centers["Mid"] - centers["Low"]), abs(centers["High"] - centers["Mid"]))
            else:
                vals = centers.dropna().to_numpy(dtype=float)
                sep_min = float(np.nanmin(np.abs(np.diff(np.sort(vals))))) if len(vals) >= 2 else np.nan
            rows.append({
                "normalization_method": method,
                "analyte": label,
                "Level": "ALL_LEVEL_SEPARATION",
                "ref_day_median_CV_%": np.nan,
                "device_day_median_CV_%": np.nan,
                "median_%bias": np.nan,
                "abs_median_%bias": np.nan,
                "IQR_%bias": np.nan,
                "min_level_center_separation": sep_min,
                "N": int(len(tmp)),
            })

    comp = pd.DataFrame(rows)
    if comp.empty:
        return comp

    metric_rows = comp[comp["Level"] != "ALL_LEVEL_SEPARATION"].copy()
    sep_rows = comp[comp["Level"] == "ALL_LEVEL_SEPARATION"].copy()

    # Raw separation is the reference for distortion. Higher separation is better.
    raw_sep = sep_rows[sep_rows["normalization_method"] == "Raw/no normalization"].set_index("analyte")["min_level_center_separation"].to_dict() if not sep_rows.empty else {}

    def _separation_ratio(r):
        denom = raw_sep.get(r["analyte"], np.nan)
        num = r.get("min_level_center_separation", np.nan)
        if not np.isfinite(denom) or denom == 0 or not np.isfinite(num):
            return np.nan
        return float(num / denom)

    sep_rows["separation_preservation_ratio"] = sep_rows.apply(_separation_ratio, axis=1)
    sep_summary = sep_rows.groupby(["normalization_method", "analyte"], as_index=False)["separation_preservation_ratio"].median()

    summary = metric_rows.groupby(["normalization_method", "analyte"], as_index=False).agg(
        mid_ref_day_CV_pct=("ref_day_median_CV_%", lambda x: np.nanmedian(x[metric_rows.loc[x.index, "Level"].astype(str).eq("Mid")]) if any(metric_rows.loc[x.index, "Level"].astype(str).eq("Mid")) else np.nan),
        all_ref_day_CV_pct=("ref_day_median_CV_%", "median"),
        all_device_day_CV_pct=("device_day_median_CV_%", "median"),
        abs_median_pctbias=("abs_median_%bias", "median"),
        iqr_pctbias=("IQR_%bias", "median"),
    )
    summary = summary.merge(sep_summary, on=["normalization_method", "analyte"], how="left")

    # Robust score: lower drift/CV/bias variation is better; preserving separation is rewarded.
    # Some methods intentionally center or z-score data, so a CV denominator can become zero and produce all-NaN
    # metrics.  Treat missing/non-finite score components as a large penalty instead of allowing idxmin() to crash.
    def _clean_metric(series: pd.Series, fallback: float = 1000.0) -> pd.Series:
        s = pd.to_numeric(series, errors="coerce").replace([np.inf, -np.inf], np.nan)
        finite = s[np.isfinite(s)]
        fill = float(np.nanmedian(finite)) if len(finite) else fallback
        return s.fillna(fill)

    mid_ref = _clean_metric(summary["mid_ref_day_CV_pct"].fillna(summary["all_ref_day_CV_pct"]))
    all_ref = _clean_metric(summary["all_ref_day_CV_pct"])
    all_dev = _clean_metric(summary["all_device_day_CV_pct"])
    bias_iqr = _clean_metric(summary["iqr_pctbias"])
    sep_ratio = pd.to_numeric(summary["separation_preservation_ratio"], errors="coerce").replace([np.inf, -np.inf], np.nan).fillna(0.0)
    sep_penalty = 100.0 * (1.0 - sep_ratio.clip(lower=0, upper=1))

    summary["normalization_score_lower_is_better"] = (
        0.40 * mid_ref +
        0.20 * all_ref +
        0.15 * all_dev +
        0.15 * bias_iqr +
        0.10 * sep_penalty
    ).replace([np.inf, -np.inf], np.nan)

    summary["recommended_for_analyte"] = False
    for label, sub in summary.groupby("analyte"):
        score = pd.to_numeric(sub["normalization_score_lower_is_better"], errors="coerce").replace([np.inf, -np.inf], np.nan)
        if score.notna().any():
            idx = score.idxmin()
            summary.loc[idx, "recommended_for_analyte"] = True

    overall = summary.groupby("normalization_method", as_index=False).agg(
        overall_score_lower_is_better=("normalization_score_lower_is_better", "median"),
        recommended_analyte_count=("recommended_for_analyte", "sum"),
    )
    overall["overall_score_lower_is_better"] = pd.to_numeric(overall["overall_score_lower_is_better"], errors="coerce").replace([np.inf, -np.inf], np.nan)
    if overall["overall_score_lower_is_better"].isna().all():
        overall["overall_score_lower_is_better"] = 1000.0
    else:
        overall["overall_score_lower_is_better"] = overall["overall_score_lower_is_better"].fillna(overall["overall_score_lower_is_better"].max() + 1000.0)
    overall = overall.sort_values(["overall_score_lower_is_better", "recommended_analyte_count"], ascending=[True, False])
    summary = summary.merge(overall, on="normalization_method", how="left")
    return summary.sort_values(["overall_score_lower_is_better", "analyte", "normalization_score_lower_is_better"], na_position="last")


# -----------------------------
# Outlier logic
# -----------------------------
def select_actual_outlier_method(method: str, vals: np.ndarray) -> str:
    if not str(method).startswith("Automatic"):
        return method
    vals = np.asarray(vals, dtype=float)
    vals = vals[np.isfinite(vals)]
    if len(vals) >= 3:
        try:
            p = stats.shapiro(vals).pvalue
            if p >= 0.05:
                return "Gcrit Grubbs-like: remove largest |value-mean|/SD if >= Gcrit"
        except Exception:
            pass
    return "Robust MAD modified-z: remove largest robust z if >= threshold"


def detect_outliers_one_group(df_group: pd.DataFrame, analyte: str, cfg: Config) -> Tuple[pd.DataFrame, pd.DataFrame]:
    work = df_group.copy()
    work[analyte] = pd.to_numeric(work[analyte], errors="coerce")
    work["is_outlier"] = False
    work["outlier_order"] = np.nan
    work["outlier_method"] = ""
    work["outlier_metric"] = np.nan
    work["outlier_threshold"] = np.nan
    work["outlier_direction"] = ""
    work["outlier_details"] = ""
    work["gcrit_mode"] = cfg.gcrit_mode
    work["gcrit_alpha"] = cfg.gcrit_alpha
    work["gcrit_tail"] = cfg.gcrit_tail

    log_rows = []
    remaining = list(work.index[work[analyte].notna()])
    method = select_actual_outlier_method(str(cfg.outlier_method), work.loc[remaining, analyte].to_numpy(dtype=float))

    if method == "None" or int(cfg.max_remove_per_group) <= 0 or len(remaining) < 3:
        return work, pd.DataFrame(log_rows)

    for step in range(int(cfg.max_remove_per_group)):
        vals = work.loc[remaining, analyte].astype(float)
        vals = vals[np.isfinite(vals)]
        if len(vals) < 3:
            break
        x = vals.to_numpy(dtype=float)
        idxs = list(vals.index)

        chosen_idx = None
        metric = np.nan
        threshold = np.nan
        direction = ""
        details = ""

        if method.startswith("Gcrit"):
            mu = float(np.mean(x))
            sd = float(np.std(x, ddof=1))
            if not np.isfinite(sd) or sd == 0:
                break
            gvals = np.abs(x - mu) / sd
            k = int(np.argmax(gvals))
            threshold = current_gcrit(len(x), cfg)
            metric = float(gvals[k])
            if np.isfinite(threshold) and metric >= threshold:
                chosen_idx = idxs[k]
                direction = "high" if x[k] > mu else "low"
                details = f"G={metric:.4g}; mean={mu:.4g}; sd={sd:.4g}; n={len(x)}; Gcrit={threshold:.4g}"
            else:
                break

        elif method.startswith("Robust MAD"):
            med = float(np.median(x))
            mad = float(np.median(np.abs(x - med)))
            if not np.isfinite(mad) or mad == 0:
                break
            modz = 0.6745 * (x - med) / mad
            k = int(np.argmax(np.abs(modz)))
            metric = float(abs(modz[k]))
            threshold = float(cfg.modified_z_threshold)
            if metric >= threshold:
                chosen_idx = idxs[k]
                direction = "high" if x[k] > med else "low"
                details = f"modified_z={modz[k]:.4g}; median={med:.4g}; MAD={mad:.4g}; threshold={threshold:.4g}"
            else:
                break

        elif method.startswith("95% robust interval"):
            med = float(np.median(x))
            rsd = robust_sd_mad(x)
            if not np.isfinite(rsd) or rsd == 0:
                break
            lo = med - float(cfg.robust_interval_z) * rsd
            hi = med + float(cfg.robust_interval_z) * rsd
            distances = np.maximum(lo - x, x - hi)
            k = int(np.argmax(distances))
            metric = float(distances[k])
            threshold = 0.0
            if metric > 0:
                chosen_idx = idxs[k]
                direction = "high" if x[k] > hi else "low"
                details = f"value outside robust interval [{lo:.4g}, {hi:.4g}]; median={med:.4g}; robust_SD={rsd:.4g}; z={cfg.robust_interval_z}"
            else:
                break

        if chosen_idx is None:
            break

        work.loc[chosen_idx, "is_outlier"] = True
        work.loc[chosen_idx, "outlier_order"] = step + 1
        work.loc[chosen_idx, "outlier_method"] = method
        work.loc[chosen_idx, "outlier_metric"] = metric
        work.loc[chosen_idx, "outlier_threshold"] = threshold
        work.loc[chosen_idx, "outlier_direction"] = direction
        work.loc[chosen_idx, "outlier_details"] = details

        row = work.loc[chosen_idx]
        log_rows.append({
            "analyte": analyte,
            "Level": row.get("Level", ""),
            "Device": row.get("Device", ""),
            "batch_id": row.get("batch_id", ""),
            "Blood Sample ID": row.get("Blood Sample ID", ""),
            "Day": row.get("Day", ""),
            "Replicate": row.get("Replicate", ""),
            "removed_order": step + 1,
            "outlier_method_requested": cfg.outlier_method,
            "outlier_method_used": method,
            "value_removed": row[analyte],
            "direction": direction,
            "outlier_metric": metric,
            "outlier_threshold": threshold,
            "details": details,
            "gcrit_mode": cfg.gcrit_mode,
            "gcrit_alpha": cfg.gcrit_alpha,
            "gcrit_tail": cfg.gcrit_tail,
        })
        remaining.remove(chosen_idx)

    return work, pd.DataFrame(log_rows)


# -----------------------------
# Validation + statistics
# -----------------------------
def validate_and_standardize(df: pd.DataFrame, analytes: List[str]) -> Tuple[bool, str]:
    df_s = standardize_base_columns(df)
    missing = [c for c in REQUIRED_BASE_COLS if c not in df_s.columns]
    if missing:
        return False, f"Missing required columns after alias handling: {missing}"
    missing_analytes = [a for a in analytes if a not in df_s.columns]
    if missing_analytes:
        return False, f"Missing analyte columns: {missing_analytes}"
    return True, "OK"


def robust_precision_components(df_clean: pd.DataFrame, analyte: str, days: List[str]) -> Tuple[float, float, float, float, float, float, float]:
    y = pd.to_numeric(df_clean[analyte], errors="coerce").dropna().to_numpy(dtype=float)
    within_sds = []
    for d in days:
        vals = pd.to_numeric(df_clean.loc[df_clean["Day"].astype(str) == str(d), analyte], errors="coerce").dropna().to_numpy(dtype=float)
        if len(vals) >= 2:
            within_sds.append(robust_sd_mad(vals))
        elif len(vals) == 1:
            within_sds.append(0.0)
    sd_repeat = float(np.median(within_sds)) if len(within_sds) else np.nan

    day_meds = []
    for d in days:
        vals = pd.to_numeric(df_clean.loc[df_clean["Day"].astype(str) == str(d), analyte], errors="coerce").dropna().to_numpy(dtype=float)
        if len(vals) > 0:
            day_meds.append(np.median(vals))
    sd_day = float(robust_sd_mad(np.array(day_meds))) if len(day_meds) >= 2 else 0.0

    sd_total = float(np.sqrt(sd_repeat**2 + sd_day**2)) if np.isfinite(sd_repeat) and np.isfinite(sd_day) else np.nan
    center = float(np.median(y)) if len(y) else np.nan
    cv_repeat = 100.0 * sd_repeat / abs(center) if np.isfinite(center) and center != 0 else np.nan
    cv_day = 100.0 * sd_day / abs(center) if np.isfinite(center) and center != 0 else np.nan
    cv_total = 100.0 * sd_total / abs(center) if np.isfinite(center) and center != 0 else np.nan
    return sd_repeat, sd_day, sd_total, center, cv_repeat, cv_day, cv_total


def bootstrap_precision_ci(df_clean: pd.DataFrame, analyte: str, days: List[str], n_boot: int, seed: int) -> Dict[str, float]:
    rng = np.random.default_rng(int(seed))
    clean = df_clean.copy()
    clean[analyte] = pd.to_numeric(clean[analyte], errors="coerce")
    clean = clean.dropna(subset=[analyte])
    if len(clean) < 3:
        return {}

    cols = ["SD_repeat", "SD_between_day", "SD_total", "CV_repeat_%", "CV_between_day_%", "CV_total_%"]
    vals = {c: [] for c in cols}
    for _ in range(int(n_boot)):
        idx = rng.choice(clean.index.to_numpy(), size=len(clean), replace=True)
        boot = clean.loc[idx].copy()
        comps = robust_precision_components(boot, analyte, days)
        row = {
            "SD_repeat": comps[0],
            "SD_between_day": comps[1],
            "SD_total": comps[2],
            "CV_repeat_%": comps[4],
            "CV_between_day_%": comps[5],
            "CV_total_%": comps[6],
        }
        for c in cols:
            if np.isfinite(row[c]):
                vals[c].append(row[c])

    out = {}
    for c in cols:
        arr = np.asarray(vals[c], dtype=float)
        arr = arr[np.isfinite(arr)]
        if len(arr):
            lo, hi = np.percentile(arr, [2.5, 97.5])
            out[f"{c}_95CI_low"] = float(lo)
            out[f"{c}_95CI_high"] = float(hi)
        else:
            out[f"{c}_95CI_low"] = np.nan
            out[f"{c}_95CI_high"] = np.nan
    return out


def compute_assumption_tests(df_clean: pd.DataFrame, analyte: str, days: List[str]) -> Dict[str, object]:
    y = pd.to_numeric(df_clean[analyte], errors="coerce").dropna().to_numpy(dtype=float)
    out = {
        "shapiro_p": np.nan,
        "normality_pass_0_05": False,
        "levene_p": np.nan,
        "variance_pass_0_05": False,
        "statistical_branch": "nonparametric_or_robust",
    }
    if len(y) >= 3:
        try:
            out["shapiro_p"] = float(stats.shapiro(y).pvalue)
            out["normality_pass_0_05"] = bool(out["shapiro_p"] >= 0.05)
        except Exception:
            pass
    groups = []
    for d in days:
        vals = pd.to_numeric(df_clean.loc[df_clean["Day"].astype(str) == str(d), analyte], errors="coerce").dropna().to_numpy(dtype=float)
        if len(vals) >= 2:
            groups.append(vals)
    if len(groups) >= 2:
        try:
            out["levene_p"] = float(stats.levene(*groups).pvalue)
            out["variance_pass_0_05"] = bool(out["levene_p"] >= 0.05)
        except Exception:
            pass
    if out["normality_pass_0_05"] and (np.isnan(out["levene_p"]) or out["variance_pass_0_05"]):
        out["statistical_branch"] = "parametric"
    return out


def compute_descriptive_ci(y: np.ndarray, normality_pass: bool, cfg: Config) -> Dict[str, float]:
    y = np.asarray(y, dtype=float)
    y = y[np.isfinite(y)]
    out = {
        "mean": float(np.mean(y)) if len(y) else np.nan,
        "sd_sample": float(np.std(y, ddof=1)) if len(y) >= 2 else np.nan,
        "median": float(np.median(y)) if len(y) else np.nan,
        "IQR": iqr(y),
        "center_95CI_low": np.nan,
        "center_95CI_high": np.nan,
        "center_CI_type": "",
    }
    if len(y) >= 2 and normality_pass:
        se = stats.sem(y, nan_policy="omit")
        tcrit = stats.t.ppf(0.975, df=len(y) - 1)
        out["center_95CI_low"] = float(out["mean"] - tcrit * se)
        out["center_95CI_high"] = float(out["mean"] + tcrit * se)
        out["center_CI_type"] = "t_95CI_for_mean"
    elif len(y) >= 2:
        lo, hi = bootstrap_ci_mean_or_median(y, use_median=True, n_boot=min(cfg.n_boot, 5000), seed=cfg.seed)
        out["center_95CI_low"] = lo
        out["center_95CI_high"] = hi
        out["center_CI_type"] = "bootstrap_95CI_for_median"
    return out


def compute_ep05_components(df_group: pd.DataFrame, analyte: str, cfg: Config, expected_n_scope: int) -> Tuple[Dict, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    out = {"analyte": analyte}
    out["Level"] = str(df_group["Level"].iloc[0]) if len(df_group) else None
    out["Device"] = str(df_group["Device"].iloc[0]) if len(df_group) else None
    out["normalization_method"] = cfg.normalization_method

    y_raw = pd.to_numeric(df_group[analyte], errors="coerce").dropna()
    out["N_raw"] = int(len(y_raw))

    df_with, outlier_log = detect_outliers_one_group(df_group, analyte, cfg)
    df_clean = df_with.loc[~df_with["is_outlier"]].copy()
    y = pd.to_numeric(df_clean[analyte], errors="coerce").dropna().to_numpy(dtype=float)

    out["N_clean"] = int(len(y))
    out["n_outliers"] = int(df_with["is_outlier"].sum())
    out["expected_n"] = int(expected_n_scope)
    out["outlier_method_requested"] = cfg.outlier_method
    used_methods = sorted([m for m in df_with.get("outlier_method", pd.Series(dtype=str)).astype(str).unique() if m])
    out["outlier_method_used"] = "; ".join(used_methods) if used_methods else ("None" if cfg.outlier_method == "None" else "No outlier removed")
    out["max_outliers_allowed"] = int(cfg.max_remove_per_group)
    out["gcrit"] = current_gcrit(len(y_raw), cfg) if cfg.outlier_method.startswith(("Gcrit", "Automatic")) else float(cfg.gcrit)
    out["gcrit_mode"] = cfg.gcrit_mode
    out["gcrit_alpha"] = cfg.gcrit_alpha
    out["gcrit_tail"] = cfg.gcrit_tail
    out["modified_z_threshold"] = cfg.modified_z_threshold
    out["robust_interval_z"] = cfg.robust_interval_z

    assumption = compute_assumption_tests(df_clean, analyte, cfg.days)
    out.update(assumption)
    out.update(compute_descriptive_ci(y, bool(out["normality_pass_0_05"]), cfg))

    out["method"] = "ROBUST_MAD"

    def robust_path():
        return robust_precision_components(df_clean, analyte, cfg.days)

    # Parametric path only if assumptions pass and the expected design is complete/no outliers.
    if out["n_outliers"] == 0 and out["N_clean"] == int(expected_n_scope) and out["statistical_branch"] == "parametric":
        try:
            df_m = df_clean.rename(columns={analyte: "value"}).copy()
            model = smf.mixedlm("value ~ 1", df_m, groups=df_m["Day"])
            res = model.fit(reml=True, method="lbfgs", disp=False)

            var_day = float(res.cov_re.iloc[0, 0]) if res.cov_re.shape == (1, 1) else 0.0
            var_within = float(res.scale)

            sd_repeat = float(np.sqrt(max(var_within, 0.0)))
            sd_day = float(np.sqrt(max(var_day, 0.0)))
            sd_total = float(np.sqrt(max(var_within + var_day, 0.0)))

            mean_val = float(np.mean(y))
            cv_repeat = 100.0 * sd_repeat / abs(mean_val) if mean_val != 0 else np.nan
            cv_day = 100.0 * sd_day / abs(mean_val) if mean_val != 0 else np.nan
            cv_total = 100.0 * sd_total / abs(mean_val) if mean_val != 0 else np.nan

            out["method"] = "MIXEDLM"
            out["center_used"] = "mean"
            out["center_value"] = mean_val
            out["SD_repeat"] = sd_repeat
            out["SD_between_day"] = sd_day
            out["SD_total"] = sd_total
            out["CV_repeat_%"] = cv_repeat
            out["CV_between_day_%"] = cv_day
            out["CV_total_%"] = cv_total

            if cfg.do_bootstrap_ci:
                out.update(bootstrap_precision_ci(df_clean, analyte, cfg.days, cfg.n_boot, cfg.seed))

            return out, df_with, df_clean, outlier_log
        except Exception:
            out["statistical_branch"] = "parametric_assumptions_passed_but_MixedLM_failed_used_robust"

    sd_repeat, sd_day, sd_total, center, cv_repeat, cv_day, cv_total = robust_path()
    out["center_used"] = "median"
    out["center_value"] = center
    out["SD_repeat"] = sd_repeat
    out["SD_between_day"] = sd_day
    out["SD_total"] = sd_total
    out["CV_repeat_%"] = cv_repeat
    out["CV_between_day_%"] = cv_day
    out["CV_total_%"] = cv_total

    if cfg.do_bootstrap_ci:
        out.update(bootstrap_precision_ci(df_clean, analyte, cfg.days, cfg.n_boot, cfg.seed))

    return out, df_with, df_clean, outlier_log


# -----------------------------
# Plotting and ZIP output
# -----------------------------
def make_histogram_png(values: np.ndarray, title: str) -> bytes:
    fig = plt.figure()
    plt.hist(values, bins=20)
    plt.title(title)
    plt.xlabel("Value")
    plt.ylabel("Count")
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


def make_boxplot_png(values: np.ndarray, title: str) -> bytes:
    fig = plt.figure()
    plt.boxplot(values, vert=True)
    plt.title(title)
    plt.ylabel("Value")
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


def _design_filter(df: pd.DataFrame, cfg: Config) -> pd.DataFrame:
    out = standardize_base_columns(df)
    if "Level" in out.columns:
        out = out[out["Level"].astype(str).isin(cfg.levels)]
    if "Day" in out.columns:
        out = out[out["Day"].astype(str).isin(cfg.days)]
    if "Device" in out.columns:
        out = out[out["Device"].astype(str).isin([str(x) for x in cfg.devices])]
    if "Replicate" in out.columns:
        reps = pd.to_numeric(out["Replicate"], errors="coerce")
        out = out[reps.isin(cfg.replicates)]
    return out.copy()


def _global_flag_audit(excluded: pd.DataFrame, cfg: Config) -> pd.DataFrame:
    cols = ["batch_id", "bloodSampleId", "analyte"]
    if excluded is None or excluded.empty:
        return pd.DataFrame(columns=cols)
    ex = _design_filter(excluded, cfg)
    rows = []
    for _, row in ex.iterrows():
        batch = row.get("batch_id", "")
        sample = row.get("Blood Sample ID", row.get("bloodSampleId", ""))
        for label, (dev_col, ref_col) in cfg.analyte_pair_map.items():
            if (dev_col in ex.columns and pd.notna(row.get(dev_col, np.nan))) or (ref_col in ex.columns and pd.notna(row.get(ref_col, np.nan))):
                rows.append({"batch_id": batch, "bloodSampleId": sample, "analyte": label})
        for analyte in cfg.analytes:
            if analyte in ex.columns and pd.notna(row.get(analyte, np.nan)):
                rows.append({"batch_id": batch, "bloodSampleId": sample, "analyte": analyte})
    return pd.DataFrame(rows, columns=cols).drop_duplicates().reset_index(drop=True)


def _format_results(summary: pd.DataFrame) -> pd.DataFrame:
    if summary is None or summary.empty:
        return pd.DataFrame()
    out = summary.copy()
    out.insert(0, "Analysis metric", out["analyte"].astype(str).str.split("__").str[1].fillna("raw") if "analyte" in out.columns else "")
    out.insert(0, "Analyte", out["analyte"].astype(str).str.split("__").str[0] if "analyte" in out.columns else "")
    preferred = [
        "Level", "Analyte", "Analysis metric", "Device", "CV_total_%", "CV_repeat_%", "CV_between_day_%",
        "shapiro_p", "normality_pass_0_05", "N_raw", "N_clean", "n_outliers", "expected_n",
        "levene_p", "variance_pass_0_05", "statistical_branch", "method", "center_used", "center_value",
        "mean", "sd_sample", "median", "IQR", "center_95CI_low", "center_95CI_high", "center_95CI_type",
        "SD_repeat", "SD_between_day", "SD_total", "normalization_method", "outlier_method_requested",
        "outlier_method_used", "max_outliers_allowed", "gcrit", "gcrit_mode", "gcrit_alpha", "gcrit_tail",
        "modified_z_threshold", "robust_interval_z",
    ]
    ordered = [c for c in preferred if c in out.columns] + [c for c in out.columns if c not in preferred and c not in {"analyte"}]
    return out[ordered]


def _format_outliers(outlier_log: pd.DataFrame, analyzed_source: pd.DataFrame, cfg: Config) -> pd.DataFrame:
    cols = ["analyte", "batch_id", "bloodSampleId", "mhs_value", "reference_value", "exclusion_reason"]
    if outlier_log is None or outlier_log.empty:
        return pd.DataFrame(columns=cols)
    src = standardize_base_columns(analyzed_source)
    rows = []
    for _, r in outlier_log.iterrows():
        metric_name = str(r.get("analyte", ""))
        label = metric_name.split("__", 1)[0]
        dev_col, ref_col = cfg.analyte_pair_map.get(label, (None, None))
        batch = r.get("batch_id", "")
        sample = r.get("Blood Sample ID", "")
        candidates = src
        if "batch_id" in candidates.columns and str(batch) != "":
            candidates = candidates[candidates["batch_id"].astype(str) == str(batch)]
        if "Blood Sample ID" in candidates.columns and str(sample) != "":
            candidates = candidates[candidates["Blood Sample ID"].astype(str) == str(sample)]
        rawrow = candidates.iloc[0] if len(candidates) else None
        mhs = rawrow.get(dev_col, np.nan) if rawrow is not None and dev_col else r.get("value_removed", np.nan)
        ref = rawrow.get(ref_col, np.nan) if rawrow is not None and ref_col else np.nan
        reason = f"{r.get('outlier_method_used', '')}: {r.get('details', '')}".strip(": ")
        rows.append({
            "analyte": label,
            "batch_id": batch,
            "bloodSampleId": sample,
            "mhs_value": mhs,
            "reference_value": ref,
            "exclusion_reason": reason,
        })
    return pd.DataFrame(rows, columns=cols)


def run_pipeline_to_excel(df: pd.DataFrame, cfg: Config) -> bytes:
    """Run the existing EP05 pipeline and return one audit-ready Excel workbook."""
    eligible_raw, excluded_global = split_global_flag_rows(df, cfg)
    df_work, analysis_analytes, pair_table = build_analysis_dataframe(eligible_raw, cfg)
    df_f = _design_filter(df_work, cfg)

    df_f["Day"] = df_f["Day"].astype(str)
    df_f["Level"] = df_f["Level"].astype(str)
    df_f["Device"] = df_f["Device"].astype(str)
    df_f["Replicate"] = pd.to_numeric(df_f["Replicate"], errors="coerce").astype(int)

    all_rows = []
    all_outlier_logs = []
    norm_comparison = evaluate_normalization_methods(df_f, cfg.analyte_pair_map, NORMALIZATION_METHODS)

    scope_tables = []
    if cfg.device_mode == "Pool all devices":
        pooled_df = df_f.copy(); pooled_df["Device"] = "pooled_all_devices"
        scope_tables.append(("pooled_all_devices", pooled_df, len(cfg.days) * len(cfg.replicates) * len(cfg.devices)))
    else:
        pooled_df = df_f.copy(); pooled_df["Device"] = "pooled_all_devices"
        scope_tables.append(("pooled_all_devices", pooled_df, len(cfg.days) * len(cfg.replicates) * len(cfg.devices)))
        for device in cfg.devices:
            scope_tables.append((str(device), df_f[df_f["Device"].astype(str) == str(device)].copy(), len(cfg.days) * len(cfg.replicates)))

    for analyte in analysis_analytes:
        for level in cfg.levels:
            for scope_name, scope_df, expected_n_scope in scope_tables:
                sub = scope_df[scope_df["Level"].astype(str) == str(level)].copy()
                if sub.empty or analyte not in sub.columns:
                    continue
                out, with_outliers, clean, outlier_log = compute_ep05_components(sub, analyte, cfg, expected_n_scope)
                all_rows.append(out)
                if outlier_log is not None and not outlier_log.empty:
                    all_outlier_logs.append(outlier_log)

    summary = pd.DataFrame(all_rows)
    outlier_log_all = pd.concat(all_outlier_logs, ignore_index=True) if all_outlier_logs else pd.DataFrame()
    results = _format_results(summary)
    outliers = _format_outliers(outlier_log_all, _design_filter(eligible_raw, cfg), cfg)
    global_log = _global_flag_audit(excluded_global, cfg)

    settings = pd.DataFrame({
        "setting": [
            "device_mode", "normalization_method", "value_output_modes", "outlier_method", "max_remove_per_group",
            "gcrit_mode", "manual_gcrit", "gcrit_alpha", "gcrit_tail", "modified_z_threshold", "robust_interval_z",
            "bootstrap_95CI", "bootstrap_iterations", "random_seed", "levels", "days", "replicates", "devices",
            "global_flag_column", "treat_all_global_flag_as_false",
        ],
        "value": [
            cfg.device_mode, cfg.normalization_method, "; ".join(cfg.value_output_modes), cfg.outlier_method, cfg.max_remove_per_group,
            cfg.gcrit_mode, cfg.gcrit, cfg.gcrit_alpha, cfg.gcrit_tail, cfg.modified_z_threshold, cfg.robust_interval_z,
            cfg.do_bootstrap_ci, cfg.n_boot, cfg.seed, ", ".join(map(str, cfg.levels)), ", ".join(map(str, cfg.days)),
            ", ".join(map(str, cfg.replicates)), ", ".join(map(str, cfg.devices)), cfg.global_flag_col or "None", cfg.treat_all_global_false,
        ],
    })

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        results.to_excel(writer, sheet_name="Results", index=False)
        outliers.to_excel(writer, sheet_name="Outliers", index=False)
        global_log.to_excel(writer, sheet_name="global flag TRUE", index=False)
        pair_table.rename(columns={"analyte": "Analyte", "device_column": "Device column", "reference_column": "Reference column"}).to_excel(writer, sheet_name="Analyte mapping", index=False)
        norm_comparison.to_excel(writer, sheet_name="Normalization comparison", index=False)
        settings.to_excel(writer, sheet_name="Settings", index=False)

        from openpyxl.styles import Font, PatternFill, Alignment
        fill = PatternFill("solid", fgColor="D9EAD3")
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            if ws.max_row and ws.max_column:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True); cell.fill = fill; cell.alignment = Alignment(vertical="center", wrap_text=True)
            for cells in ws.columns:
                letter = cells[0].column_letter
                width = max((len(str(c.value)) if c.value is not None else 0 for c in cells[:min(ws.max_row, 200)]), default=8) + 2
                ws.column_dimensions[letter].width = min(max(width, 10), 42)
    return buf.getvalue()


# -----------------------------
# Streamlit UI
# -----------------------------
st.set_page_config(page_title="EP05 Precision App", layout="wide")
st.title("EP05 Long-Term Precision + Reference Drift Normalization")

st.markdown(
    """
**Steps**
1) Upload the XLSX  
2) Confirm/adjust Levels, Days, Replicates, Devices  
3) Select paired CBC analytes and compare normalization methods  
4) Choose the recommended method or manually select another  
5) Choose pooled/per-device handling and outlier settings  
6) Run → download one Excel workbook
"""
)

uploaded = st.file_uploader("Upload the XLSX", type=["xlsx"])
if uploaded is None:
    st.info("Upload your EP05-style XLSX to begin.")
    st.stop()

try:
    df_raw = pd.read_excel(uploaded, engine="openpyxl")
    df = standardize_base_columns(df_raw)
except Exception as e:
    st.error(f"Could not read Excel file: {e}")
    st.stop()

st.subheader("Preview")
st.dataframe(df.head(20), use_container_width=True)


def present_values(col: str, defaults: List[str]) -> Tuple[List[str], List[str]]:
    if col in df.columns:
        vals = sorted(list({str(x) for x in df[col].dropna().tolist()}))
        if len(defaults) == 0:
            return vals, vals
        if all(str(d) in vals for d in defaults):
            return [str(d) for d in defaults], vals
        return vals, vals
    return [str(d) for d in defaults], [str(d) for d in defaults]


def present_int_values(col: str, defaults: List[int]) -> Tuple[List[int], List[int]]:
    if col in df.columns:
        vals = sorted(list({int(x) for x in df[col].dropna().tolist()}))
        if all(int(d) in vals for d in defaults):
            return defaults, vals
        return vals, vals
    return defaults, defaults


levels_default, levels_all = present_values("Level", DEFAULT_LEVELS)
days_default, days_all = present_values("Day", DEFAULT_DAYS)
devices_default, devices_all = present_values("Device", DEFAULT_DEVICES)
rep_default, rep_all = present_int_values("Replicate", DEFAULT_REPLICATES)

st.subheader("Design selection")
c1, c2, c3, c4 = st.columns(4)
with c1:
    levels = st.multiselect("Levels", options=levels_all, default=levels_default)
with c2:
    days = st.multiselect("Days", options=days_all, default=days_default)
with c3:
    replicates = st.multiselect("Replicates", options=rep_all, default=rep_default)
with c4:
    devices = st.multiselect("Devices", options=devices_all, default=devices_default)

flag_options = ["None"] + list(df.columns)
global_guess = next((c for c in df.columns if c.lower() == "global_flag"), None)
g1, g2 = st.columns([2, 3])
with g1:
    global_flag_col = st.selectbox(
        "Global flag column", flag_options,
        index=flag_options.index(global_guess) if global_guess in flag_options else 0,
    )
with g2:
    treat_all_global_false = st.checkbox(
        "Treat all rows as global_flag = FALSE when no flag column is selected", value=False
    )

if global_flag_col != "None" and not treat_all_global_false:
    _ui_flag_mask = df[global_flag_col].map(normalize_bool).fillna(False).astype(bool)
    df_eligible_ui = df.loc[~_ui_flag_mask].copy()
else:
    df_eligible_ui = df.copy()

st.subheader("Paired analytes and normalization")
mapping_seed_all = default_analyte_mapping_table(df_eligible_ui)

# Explicit analyte/model selector, matching the Short-Term app behavior.
# Options and selected tags are shown as uploaded device-result column names.
if not mapping_seed_all.empty:
    detected_model_cols = mapping_seed_all["Device column"].astype(str).tolist()
else:
    detected_model_cols = []
preferred_present = [c for c in DEFAULT_ANALYTE_COLUMN_ORDER if c in detected_model_cols]
ordered_model_options = preferred_present + [c for c in detected_model_cols if c not in preferred_present]
selected_model_cols = st.multiselect(
    "Analyte/model columns to analyze",
    options=ordered_model_options,
    default=ordered_model_options,
    help="Choose exactly which analyte/model result columns to include. The suggested order matches the Short-Term app.",
)
mapping_seed = mapping_seed_all[mapping_seed_all["Device column"].astype(str).isin(selected_model_cols)].copy()
if not mapping_seed.empty:
    selected_order = {name: i for i, name in enumerate(selected_model_cols)}
    mapping_seed["__order"] = mapping_seed["Device column"].astype(str).map(selected_order)
    mapping_seed = mapping_seed.sort_values("__order").drop(columns="__order").reset_index(drop=True)

st.caption("Edit this table to add or remove analyzer models. Multiple models may share the same reference column (for example PLT, PLT 2 and PLT 3 can all use PLT_ref).")
device_measure_options = [c for c in df_eligible_ui.columns if pd.to_numeric(df_eligible_ui[c], errors="coerce").notna().sum() >= 3]
reference_options = [c for c in df_eligible_ui.columns if c.lower().endswith("_ref")]
mapping_table = st.data_editor(
    mapping_seed,
    num_rows="dynamic",
    use_container_width=True,
    hide_index=True,
    column_config={
        "Include": st.column_config.CheckboxColumn("Include", default=True),
        "Analyte": st.column_config.TextColumn("Analyte", required=True),
        "Device column": st.column_config.SelectboxColumn("Device column", options=device_measure_options, required=True),
        "Reference column": st.column_config.SelectboxColumn("Reference column", options=reference_options, required=True),
    },
    key="analyte_mapping_editor",
)
selected_pair_map = parse_analyte_mapping_table(mapping_table, df_eligible_ui)
paired_analytes = list(selected_pair_map.keys())
if selected_pair_map:
    st.dataframe(pd.DataFrame([{"Analyte": k, "Device column": v[0], "Reference column": v[1]} for k, v in selected_pair_map.items()]), use_container_width=True)
else:
    st.warning("Select or add at least one valid analyte mapping, or choose a raw analyte below.")


compare_now = st.checkbox("Show normalization comparison and recommendation", value=True)
recommended_method = "Raw/no normalization"
norm_comparison = pd.DataFrame()
if compare_now and selected_pair_map:
    with st.spinner("Comparing normalization methods..."):
        norm_comparison = evaluate_normalization_methods(df_eligible_ui[df_eligible_ui["Level"].astype(str).isin(levels)] if levels else df_eligible_ui, selected_pair_map, NORMALIZATION_METHODS)
    if not norm_comparison.empty:
        recommended_method = str(norm_comparison.sort_values("overall_score_lower_is_better").iloc[0]["normalization_method"])
        st.success(f"Recommended normalization for this dataset: {recommended_method}")
        st.dataframe(norm_comparison, use_container_width=True)
    else:
        st.info("Normalization comparison could not be computed; using raw/no normalization by default.")

norm_default_index = NORMALIZATION_METHODS.index(recommended_method) if recommended_method in NORMALIZATION_METHODS else 0
normalization_method = st.selectbox(
    "Normalization method to use for the final EP05 run",
    NORMALIZATION_METHODS,
    index=norm_default_index,
)

value_output_modes = st.multiselect(
    "Metrics to calculate/analyze from each device-reference pair",
    VALUE_OUTPUT_MODES,
    default=["Device normalized value", "%Bias: 100*(device-reference)/reference"],
)

st.subheader("Additional raw analyte columns, optional")
observed_candidate_cols = [c for c in df_eligible_ui.columns if c not in REQUIRED_BASE_COLS]
extra_analytes = st.multiselect(
    "Optional: also analyze raw columns directly without paired normalization",
    options=sorted(observed_candidate_cols),
    default=[],
)

st.subheader("Analysis settings")
c1, c2, c3, c4 = st.columns(4)
with c1:
    device_mode = st.selectbox("Device handling", ["Pool all devices", "Analyze each device separately + pooled"])
with c2:
    do_bootstrap_ci = st.checkbox("Bootstrap 95% CIs", value=False)
with c3:
    n_boot = st.number_input("Bootstrap iterations", min_value=200, max_value=20000, value=2000, step=200)
with c4:
    seed = st.number_input("Random seed", min_value=1, max_value=999999, value=123, step=1)

st.subheader("Optional outlier detection")
outlier_method = st.selectbox(
    "Outlier method for cleaned/sensitivity results",
    [
        "None",
        "Automatic: Grubbs if Shapiro normal, otherwise Robust MAD",
        "Gcrit Grubbs-like: remove largest |value-mean|/SD if >= Gcrit",
        "Robust MAD modified-z: remove largest robust z if >= threshold",
        "95% robust interval: remove most extreme outside median ± z*MAD_SD",
    ],
    index=1,
)
c1, c2, c3, c4 = st.columns(4)
with c1:
    max_remove_per_group = st.selectbox("Max outliers to remove per analyte/level/scope", [0, 1, 2], index=1)
with c2:
    gcrit_mode = st.selectbox("Gcrit mode", ["Manual Gcrit value", "Automatic from n, alpha, tail"], index=1)
with c3:
    gcrit = st.number_input("Manual Gcrit value", min_value=0.0, value=3.135, step=0.001, format="%.3f")
with c4:
    gcrit_alpha = st.number_input("Automatic Gcrit alpha", min_value=0.0001, max_value=0.2, value=0.01, step=0.001, format="%.4f")

c1, c2, c3 = st.columns(3)
with c1:
    gcrit_tail = st.selectbox("Automatic Gcrit tail", ["Two-sided", "One-sided"], index=0)
with c2:
    modified_z_threshold = st.number_input("MAD modified-z threshold", min_value=0.1, value=3.5, step=0.1)
with c3:
    robust_interval_z = st.number_input("Robust interval z", min_value=0.5, value=1.96, step=0.01)

expected_n_single_device = len(days) * len(replicates)
expected_n_pooled = len(days) * len(replicates) * len(devices)
st.caption(
    f"Expected N per single device group = Days × Replicates = {len(days)} × {len(replicates)} = **{expected_n_single_device}**. "
    f"Pooled expected N across selected devices = {len(days)} × {len(replicates)} × {len(devices)} = **{expected_n_pooled}**."
)
st.markdown("""
**Assumption-aware testing:** each analyte/level/scope now receives Shapiro-Wilk normality testing and Levene variance testing.  
If assumptions pass, the app attempts the parametric MixedLM precision estimate. If assumptions fail or the model cannot fit, it uses robust MAD-based precision.  
**Outlier default:** automatic outlier handling uses Grubbs only when Shapiro-Wilk supports normality; otherwise it uses robust MAD.
""")

# Validation is run after normalized columns are generated.
placeholder_cfg = Config(
    analytes=extra_analytes,
    levels=levels,
    days=days,
    replicates=[int(r) for r in replicates],
    devices=devices,
    gcrit=float(gcrit),
    expected_n=int(expected_n_single_device),
    device_mode=device_mode,
    outlier_method=outlier_method,
    max_remove_per_group=int(max_remove_per_group),
    gcrit_mode=gcrit_mode,
    gcrit_alpha=float(gcrit_alpha),
    gcrit_tail=gcrit_tail,
    modified_z_threshold=float(modified_z_threshold),
    robust_interval_z=float(robust_interval_z),
    do_bootstrap_ci=bool(do_bootstrap_ci),
    n_boot=int(n_boot),
    seed=int(seed),
    paired_analytes=paired_analytes,
    normalization_method=normalization_method,
    value_output_modes=value_output_modes,
    analyte_pair_map=selected_pair_map,
)
try:
    df_check, analysis_analytes_check, _ = build_analysis_dataframe(df_eligible_ui, placeholder_cfg)
    ok, msg = validate_and_standardize(df_check, analysis_analytes_check)
except Exception as e:
    ok, msg = False, str(e)

if not ok:
    st.error(msg)
    st.stop()

run_btn = st.button("Run EP05 analysis", type="primary")
if run_btn:
    if len(levels) == 0 or len(days) == 0 or len(replicates) == 0 or len(devices) == 0:
        st.error("Please select at least one Level, Day, Replicate, and Device.")
        st.stop()
    if len(selected_pair_map) == 0 and len(extra_analytes) == 0:
        st.error("Please select at least one paired analyte or raw analyte column.")
        st.stop()
    if len(value_output_modes) == 0 and len(selected_pair_map) > 0:
        st.error("Please select at least one paired metric to analyze.")
        st.stop()

    cfg = Config(
        analytes=extra_analytes,
        levels=levels,
        days=days,
        replicates=[int(r) for r in replicates],
        devices=devices,
        gcrit=float(gcrit),
        expected_n=int(expected_n_single_device),
        device_mode=device_mode,
        outlier_method=outlier_method,
        max_remove_per_group=int(max_remove_per_group),
        gcrit_mode=gcrit_mode,
        gcrit_alpha=float(gcrit_alpha),
        gcrit_tail=gcrit_tail,
        modified_z_threshold=float(modified_z_threshold),
        robust_interval_z=float(robust_interval_z),
        do_bootstrap_ci=bool(do_bootstrap_ci),
        n_boot=int(n_boot),
        seed=int(seed),
        paired_analytes=paired_analytes,
        normalization_method=normalization_method,
        value_output_modes=value_output_modes,
        analyte_pair_map=selected_pair_map,
        global_flag_col=None if global_flag_col == "None" else global_flag_col,
        treat_all_global_false=bool(treat_all_global_false),
    )

    if (not cfg.treat_all_global_false) and (not cfg.global_flag_col):
        st.error("Select a Global flag column, or tick the option to treat all rows as global_flag = FALSE.")
        st.stop()

    with st.spinner("Running analysis, excluding global_flag=TRUE rows, and creating the combined workbook..."):
        excel_bytes = run_pipeline_to_excel(df, cfg)

    st.success("Done. All results and audit tables are in one Excel workbook.")
    st.download_button(
        label="Download combined Excel results",
        data=excel_bytes,
        file_name="imprecision_long_results.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
